"""Build all editable/printable paper formats from one manuscript and data source."""
from pathlib import Path
import json,subprocess,re,shutil,os,tempfile
from copy import deepcopy
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent;P=ROOT/'paper';BASE=P/'第一二问完整论文'
from manuscript import render_manuscript, publish_repository_paper, split_chapters
FONT_SERIF='Times New Roman' if os.name=='nt' else 'Liberation Serif'
FONT_CJK='SimSun' if os.name=='nt' else 'Noto Serif CJK SC'
FONT_SANS='Microsoft YaHei' if os.name=='nt' else 'Noto Sans CJK SC'

def styles(doc):
    for sec in doc.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.25);sec.right_margin=Cm(2.25)
        sec.header_distance=Cm(.85);sec.footer_distance=Cm(.85)
        h=sec.header.paragraphs[0];h.text='第一、二问：完整模型、全局证明与解析策略';h.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in h.runs:r.font.size=Pt(8);r.font.name=FONT_SANS
        f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    for st in doc.styles:
        if st.type==1:
            st.font.name=FONT_SERIF;st.font.size=Pt(10.5);st.font.color.rgb=RGBColor(0,0,0)
            rpr=st.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
            if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
            fonts.set(qn('w:eastAsia'),FONT_CJK)
            st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(5)
    for name,size in [('Heading 1',18),('Heading 2',13),('Heading 3',11.5),('Heading 4',10.5)]:
        st=next(st for st in doc.styles if st.style_id==name.replace(' ',''));st.font.size=Pt(size);st.font.bold=True
        st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.keep_with_next=True;st.paragraph_format.keep_together=True
    for name in ['Caption','Image Caption','Table Caption']:
        if any(st.style_id==name.replace(' ','') for st in doc.styles):
            st=next(st for st in doc.styles if st.style_id==name.replace(' ',''));st.font.size=Pt(9);st.font.italic=False;st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
            st.paragraph_format.keep_with_next=False
    for p in doc.paragraphs:
        if p._p.xpath('.//w:drawing'):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        if p._p.xpath('.//m:oMathPara'):
            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7);p.paragraph_format.keep_together=True
    if doc.paragraphs:
        first=doc.paragraphs[0];first.alignment=WD_ALIGN_PARAGRAPH.CENTER;first.paragraph_format.space_before=Pt(0);first.paragraph_format.space_after=Pt(12)
    for tb in doc.tables:
        tb.autofit=True
        for i,row in enumerate(tb.rows):
            pr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
            if i==0:repeat=OxmlElement('w:tblHeader');pr.append(repeat)
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.line_spacing=1.1;para.paragraph_format.space_after=Pt(3)
                    for run in para.runs:run.font.size=Pt(9);run.font.bold=(i==0)
    if doc.tables:
        # Compact the proof-to-code map enough to retain references on the
        # same final page, without shrinking the main mathematical exposition.
        for row in doc.tables[-1].rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing=1.0
                    paragraph.paragraph_format.space_after=Pt(0)
    in_references=False
    for paragraph in doc.paragraphs:
        if paragraph.text=='参考文献':
            in_references=True;paragraph.paragraph_format.space_before=Pt(7)
        elif in_references:
            paragraph.paragraph_format.line_spacing=1.0;paragraph.paragraph_format.space_after=Pt(3)
            for r in paragraph.runs:r.font.size=Pt(9.5)
    doc.core_properties.title='有界测向误差下的交会定位与第二检测点极小极大选择'
    doc.core_properties.author='';doc.core_properties.subject='第一、二问完整解答与证明'

def main():
    q1=json.loads((ROOT/'results/q1_summary.json').read_text());q2=json.loads((ROOT/'results/q2_summary.json').read_text())
    text=(P/'正文模板.md').read_text(encoding='utf-8')
    t1='\n'.join(f"| {r['n']} | {r['bounded']}/{r['total']} | {r['median_D']:.3f} | {r['p25_D']:.3f}—{r['p75_D']:.3f} |" for r in q1['summary'])
    labs={'minimax':'极小极大解析策略','along_750':'沿示向度前进750米','perpendicular_1000':'垂线方向1000米','random_omega':'Ω内随机点'}
    t2='\n'.join(f"| {labs[r['strategy']]} | {r['N']} | {r['no_signal']} | {r['mean_D_among_valid']:.3f} | {r['p95_D_among_valid']:.3f} |" for r in q2['summary'] if r['pose']=='all')
    text=render_manuscript(text.replace('{{Q1_TABLE}}',t1).replace('{{Q2_TABLE}}',t2),q1,q2)
    BASE.with_suffix('.md').write_text(text,encoding='utf-8')
    # Separate replaceable chapters retain the same derivations, not abbreviated rewrites.
    split_chapters(text)
    if not shutil.which('pandoc'):
        import pypandoc
        os.environ['PATH']=str(Path(pypandoc.get_pandoc_path()).parent)+os.pathsep+os.environ['PATH']
    if os.name=='nt' and not shutil.which('xelatex'):
        candidate=Path(os.environ['LOCALAPPDATA'])/'Programs/MiKTeX/miktex/bin/x64'
        if (candidate/'xelatex.exe').exists():os.environ['PATH']=str(candidate)+os.pathsep+os.environ['PATH']
    log=[]
    def run(args):
        z=subprocess.run(args,cwd=P,text=True,encoding='utf-8',errors='replace',capture_output=True);log.append(' '.join(map(str,args))+'\nexit_code='+str(z.returncode)+'\n'+z.stdout+'\n'+z.stderr)
        if z.returncode:
            (ROOT/'logs/build_paper.txt').write_text('\n'.join(log),encoding='utf-8')
            raise RuntimeError(z.stdout[-5000:]+'\n'+z.stderr)
    # Word equation numbers are added as OMML below. Strip only tags from the
    # conversion input: recent Pandoc rejects an alphanumeric tag in equation A7.
    with tempfile.TemporaryDirectory(prefix='q12-word-') as staging:
        word_input=Path(staging)/'word.md'
        word_input.write_text(re.sub(r'\\tag\{[^}]+\}','',text),encoding='utf-8')
        run(['pandoc',str(word_input),'-f','markdown+tex_math_dollars+raw_tex',
             '--resource-path='+str(P),'-o',BASE.with_suffix('.docx').name])
    d=Document(BASE.with_suffix('.docx'));styles(d)
    blocks=re.findall(r'\$\$(.*?)\$\$',text,re.S)
    displays=d._element.xpath('.//m:oMathPara')
    if len(displays)!=len(blocks):raise RuntimeError('Equation count mismatch in DOCX conversion')
    for block,disp in zip(blocks,displays):
        tag=re.search(r'\\tag\{([^}]+)\}',block)
        if tag:
            math=disp.find(qn('m:oMath'))
            spacer=OxmlElement('m:r');sp=OxmlElement('m:t');sp.text='\u2003';spacer.append(sp);math.append(spacer)
            delim=OxmlElement('m:d');pr=OxmlElement('m:dPr')
            for name,value in [('begChr','('),('endChr',')'),('sepChr','')]:
                node=OxmlElement('m:'+name);node.set(qn('m:val'),value);pr.append(node)
            delim.append(pr);e=OxmlElement('m:e');mr=OxmlElement('m:r');rpr=OxmlElement('m:rPr')
            sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');rpr.append(sty);mr.append(rpr)
            t=OxmlElement('m:t');t.text=tag.group(1);mr.append(t);e.append(mr);delim.append(e);math.append(delim)
    if any('$$' in p.text or '\\rm' in p.text for p in d.paragraphs):raise RuntimeError('Unconverted TeX in Word')
    # LibreOffice's OMML import substitutes ')' for several growing right
    # delimiters. Use native, literal delimiters for non-parenthesis pairs;
    # formulas remain editable OMML, not images.
    for delim in list(d._element.xpath('.//m:d'))[::-1]:
        dp=delim.find(qn('m:dPr'))
        if dp is None:continue
        bc=dp.find(qn('m:begChr'));ec=dp.find(qn('m:endChr'))
        if bc is None or ec is None:continue
        left,right=bc.get(qn('m:val')),ec.get(qn('m:val'))
        if (left,right)==('(',')'):continue
        if len(delim.findall(qn('m:e')))!=1:continue
        parent=delim.getparent();index=parent.index(delim);pieces=[]
        for value in [left,None,right]:
            if value is None:
                pieces.extend(deepcopy(child) for child in delim.find(qn('m:e')))
            elif value:
                r=OxmlElement('m:r');pr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');pr.append(sty);r.append(pr)
                t=OxmlElement('m:t');t.text='∣' if value=='|' else value;r.append(t);pieces.append(r)
        parent.remove(delim)
        for offset,node in enumerate(pieces):parent.insert(index+offset,node)
    for t in d._element.xpath('.//m:t'):
        if t.text=='|':t.text='∣'
    d.save(BASE.with_suffix('.docx'))
    header=r'''\usepackage{fancyhdr}
\pagestyle{fancy}\fancyhf{}
\fancyhead[C]{\small 第一、二问：完整模型、全局证明与解析策略}
\fancyfoot[C]{\thepage}
\setlength{\headheight}{15pt}
\setlength{\parindent}{0pt}
\setlength{\parskip}{0.35em}
\usepackage{float}
\floatplacement{figure}{H}
\usepackage{caption}
\captionsetup{font=small,labelformat=empty}
\setlength{\emergencystretch}{3em}
'''
    (P/'header.tex').write_text(header,encoding='utf-8')
    run(['pandoc',BASE.with_suffix('.md').name,'--standalone','--pdf-engine=xelatex',
         '-V','documentclass=ctexart','-V','classoption=fontset=none','-V','fontsize=11pt',
         '-V','mainfont='+FONT_SERIF,'-V','CJKmainfont='+FONT_CJK,'-V','CJKsansfont='+FONT_SANS,'-V','CJKmonofont='+FONT_SANS,
         '-V','geometry:margin=2.1cm','-V','linestretch=1.12','-H','header.tex','-o',BASE.with_suffix('.tex').name])
    if not shutil.which('xelatex'):raise RuntimeError('Install XeLaTeX to rebuild the PDF')
    for _ in range(2):run(['xelatex','-interaction=nonstopmode','-halt-on-error',BASE.with_suffix('.tex').name])
    (ROOT/'logs/build_paper.txt').write_text('\n'.join(log),encoding='utf-8')
    for ext in ['.aux','.log','.out']:BASE.with_suffix(ext).unlink(missing_ok=True)
    publish_repository_paper()
    print('Paper built: PDF, DOCX, Markdown, LaTeX, and separate Q1/Q2 chapters.')

if __name__=='__main__':main()
